import io
import unittest
from datetime import UTC, datetime

from openpyxl import Workbook

from autonomous_mes.application.turning_plan_import import preview_turning_plan_workbook


def turning_workbook() -> bytes:
    workbook = Workbook()
    plan = workbook.active
    plan.title = "车工每日计划"
    plan.append(["人员", "物料编码", "名称", "数量", "计划完成", "状态", "单件工时"])
    plan.append(["王超伟", "MAT-OLD", "已完成本体", 1, "2026-09-15", "9.12完成", 20])
    plan.append(["王超伟", "MAT-NEW", "新本体", 2, "2026-09-20", "", 35])
    plan.append([None, "MAT-NEXT", "新底座", 1, "2026-09-21", "", 15])
    skills = workbook.create_sheet("专线列表及工时")
    skills.append(["人员", "加工种类", "8小时工时", "加班3小时工时"])
    skills.append(["王超伟", "本体", 160, 180])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class TurningPlanImportTests(unittest.TestCase):
    def test_maps_unfinished_rows_and_embedded_capacity_to_snapshot(self) -> None:
        result = preview_turning_plan_workbook(
            turning_workbook(), "车工计划.xlsx", "WS-MACH-01",
            datetime(2026, 9, 13, tzinfo=UTC),
        )

        assert result is not None
        self.assertTrue(result["valid"], result["issues"])
        self.assertEqual("TURNING_PLAN_ADAPTER", result["mappingMode"])
        self.assertEqual(2, result["stats"]["workOrderCount"])
        self.assertEqual(1, result["stats"]["completedExcludedCount"])
        self.assertEqual(1, result["stats"]["resourceCount"])
        order = result["snapshot"]["workOrders"][0]
        self.assertEqual("MAT-NEW", order["productRevisionId"].removeprefix("MATERIAL-"))
        self.assertEqual("TURN", order["operations"][0]["operationCode"])
        self.assertEqual(35, order["operations"][0]["minutesPerUnit"])
        self.assertIsNotNone(order["operations"][0]["assignedResourceId"])

    def test_prefers_daily_plan_over_earlier_matching_sheet(self) -> None:
        workbook = Workbook()
        drafting = workbook.active
        drafting.title = "工艺编制中"
        headers = ["人员", "物料编码", "名称", "数量", "计划完成", "状态", "单件工时"]
        drafting.append(headers)
        drafting.append(["王超伟", "MAT-DRAFT", "编制项", 9, "2026-09-22", "", 90])
        daily = workbook.create_sheet("车工每日计划")
        daily.append(headers)
        daily.append(["王超伟", "MAT-DAILY", "日计划", 1, "2026-09-20", "", 20])
        skills = workbook.create_sheet("专线列表及工时")
        skills.append(["人员", "加工种类", "8小时工时", "加班3小时工时"])
        skills.append(["王超伟", "本体", 160, 180])
        output = io.BytesIO()
        workbook.save(output)

        result = preview_turning_plan_workbook(
            output.getvalue(), "车工计划.xlsx", "WS-MACH-01"
        )

        assert result is not None
        self.assertEqual(1, result["stats"]["workOrderCount"])
        order = result["snapshot"]["workOrders"][0]
        self.assertEqual("MATERIAL-MAT-DAILY", order["productRevisionId"])

    def test_returns_none_for_unrelated_workbook(self) -> None:
        workbook = Workbook()
        workbook.active.append(["其他", "字段"])
        output = io.BytesIO()
        workbook.save(output)

        self.assertIsNone(
            preview_turning_plan_workbook(output.getvalue(), "other.xlsx", "WS-1")
        )
