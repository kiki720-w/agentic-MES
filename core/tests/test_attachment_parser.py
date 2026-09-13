import io
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

from autonomous_mes.application.attachment_parser import parse_attachment
from autonomous_mes.domain.errors import ValidationError


class AttachmentParserTests(unittest.TestCase):
    def test_text_and_json_are_extracted_locally(self) -> None:
        text = parse_attachment("设备 M-08 报警 E42".encode(), "alarm.txt")
        structured = parse_attachment(b'{"workOrder":"WO-18","quantity":12}', "order.json")

        self.assertEqual("PARSED", text["status"])
        self.assertIn("M-08", text["text"])
        self.assertEqual("LOCAL_TEXT", text["parser"])
        self.assertIn('"quantity": 12', structured["text"])
        self.assertTrue(text["localOnly"])

    def test_xlsx_extracts_sheet_cells(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "工单"
        sheet.append(["工单号", "数量"])
        sheet.append(["WO-900", 24])
        content = io.BytesIO()
        workbook.save(content)

        result = parse_attachment(content.getvalue(), "schedule.xlsx")

        self.assertEqual("WORKBOOK", result["kind"])
        self.assertIn("WO-900\t24", result["text"])
        self.assertEqual(2, result["metadata"]["rowCount"])
        self.assertEqual("Docling", result["metadata"]["engine"])
        self.assertIn("工单", result["summary"])

    def test_xlsx_preserves_every_sheet_in_summary(self) -> None:
        workbook = Workbook()
        first = workbook.active
        first.title = "每日计划"
        first.append(["人员", "物料编码", "数量"])
        first.append(["王师傅", "MAT-1", 4])
        capacity = workbook.create_sheet("人员能力")
        capacity.append(["人员", "加工种类", "8小时工时"])
        capacity.append(["王师傅", "轴类", 160])
        content = io.BytesIO()
        workbook.save(content)

        result = parse_attachment(content.getvalue(), "车工计划.xlsx")

        self.assertIn("每日计划", result["summary"])
        self.assertIn("人员能力", result["summary"])
        self.assertIn("人员=王师傅", result["summary"])
        self.assertEqual("REQUIRES_FIELD_MAPPING", result["metadata"]["mappingStatus"])
        self.assertEqual(2, result["metadata"]["sheetCount"])

    def test_docx_extracts_paragraphs_and_tables(self) -> None:
        document = Document()
        document.add_paragraph("换刀后复测首件")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "设备"
        table.cell(0, 1).text = "M-12"
        content = io.BytesIO()
        document.save(content)

        result = parse_attachment(content.getvalue(), "instruction.docx")

        self.assertEqual("DOCUMENT", result["kind"])
        self.assertIn("换刀后复测首件", result["text"])
        self.assertIn("设备\tM-12", result["text"])

    @unittest.skipUnless(Path("C:/Windows/Fonts/arial.ttf").exists(), "Windows font unavailable")
    def test_image_ocr_extracts_visible_text(self) -> None:
        image = Image.new("RGB", (1200, 260), "white")
        font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 72)
        ImageDraw.Draw(image).text((35, 75), "CAPAXION ORDER 48271", fill="black", font=font)
        content = io.BytesIO()
        image.save(content, format="PNG")

        result = parse_attachment(content.getvalue(), "order.png")

        self.assertEqual("IMAGE_OCR", result["kind"])
        self.assertEqual("PARSED", result["status"])
        self.assertIn("48271", result["text"])
        self.assertGreater(result["metadata"]["ocrMeanConfidence"], 0.8)

    @unittest.skipUnless(Path("C:/Windows/Fonts/arial.ttf").exists(), "Windows font unavailable")
    def test_scanned_pdf_uses_embedded_image_ocr(self) -> None:
        image = Image.new("RGB", (1200, 260), "white")
        font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 72)
        ImageDraw.Draw(image).text((35, 75), "PDF ORDER 72819", fill="black", font=font)
        content = io.BytesIO()
        image.save(content, format="PDF")

        result = parse_attachment(content.getvalue(), "scan.pdf")

        self.assertEqual("DOCUMENT", result["kind"])
        self.assertEqual("PARSED", result["status"])
        self.assertIn("72819", result["text"])
        self.assertEqual(1, result["metadata"]["ocrPageCount"])

    def test_rejects_executable_and_oversized_content(self) -> None:
        with self.assertRaises(ValidationError):
            parse_attachment(b"MZ", "payload.exe")
