"""Adapter for the common Chinese turning-plan workbook used by CAPAXION pilots."""

from __future__ import annotations

import io
from datetime import UTC, date, datetime, time
from hashlib import sha256
from pathlib import Path
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]

from .agent_workbook_import import preview_capacity_workbook
from .spreadsheet_import import snapshot_fingerprint

MAX_IMPORT_ROWS = 5_000
PLAN_HEADERS = {
    "人员", "物料编码", "名称", "数量", "计划完成", "状态", "单件工时",
}


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _due_at(value: Any) -> str | None:
    parsed: datetime | None = None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time(23, 59))
    else:
        text = _clean(value)
        if text:
            try:
                parsed = datetime.fromisoformat(text)
            except ValueError:
                try:
                    parsed = datetime.strptime(text, "%Y/%m/%d").replace(tzinfo=UTC)
                except ValueError:
                    return None
    if parsed is None:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=UTC)
    return parsed.isoformat()


def _number(value: Any) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if number > 0 else None


def preview_turning_plan_workbook(
    content: bytes,
    filename: str,
    workshop_id: str,
    observed_at: datetime | None = None,
) -> dict[str, Any] | None:
    """Return a standard scheduling preview, or None when this is another workbook."""
    if not filename.lower().endswith(".xlsx"):
        return None
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - the standard parser owns invalid-workbook errors
        return None
    candidates: list[tuple[int, Any, int, dict[str, int]]] = []
    for sheet in workbook.worksheets[:20]:
        for row_number, row in enumerate(
            sheet.iter_rows(
                min_row=1,
                max_row=min(30, sheet.max_row),
                max_col=min(64, sheet.max_column),
                values_only=True,
            ),
            start=1,
        ):
            fields = {_clean(value): index for index, value in enumerate(row) if _clean(value)}
            if PLAN_HEADERS.issubset(fields):
                title = _clean(sheet.title)
                priority = 0 if "每日计划" in title else 1 if "排产" in title else 2
                candidates.append((priority, sheet, row_number, fields))
                break
    if not candidates:
        return None
    _, plan_sheet, header_row, indexes = min(candidates, key=lambda item: item[0])

    digest = sha256(content).hexdigest()
    capacity = preview_capacity_workbook(content, filename, workshop_id, [])
    resources = []
    person_codes: dict[str, str] = {}
    for action in capacity.get("actions", []):
        source = action["resource"]
        person_codes[str(source["name"])] = str(source["code"])
        resources.append({
            "externalId": source["code"],
            "code": source["code"],
            "name": source["name"],
            "resourceType": "PERSON",
            "workCenterId": source["workCenterId"],
            "dailyCapacityMinutes": source["dailyCapacityMinutes"],
            "overtimeCapacityMinutes": source["overtimeCapacityMinutes"],
            "capabilityCodes": source["capabilityCodes"],
            "active": True,
            "state": "UNKNOWN",
        })

    orders: list[dict[str, Any]] = []
    current_person = ""
    completed_count = 0
    invalid_count = 0
    for row_number, row in enumerate(
        plan_sheet.iter_rows(
            min_row=header_row + 1,
            max_row=min(plan_sheet.max_row, header_row + MAX_IMPORT_ROWS),
            max_col=min(64, plan_sheet.max_column),
            values_only=True,
        ),
        start=header_row + 1,
    ):
        person = _clean(row[indexes["人员"]])
        if person:
            current_person = person
        material = _clean(row[indexes["物料编码"]])
        product_name = _clean(row[indexes["名称"]])
        quantity_value = _number(row[indexes["数量"]])
        due = _due_at(row[indexes["计划完成"]])
        minutes = _number(row[indexes["单件工时"]])
        status = _clean(row[indexes["状态"]])
        if not material and not product_name:
            continue
        if "完成" in status or "已完" in status:
            completed_count += 1
            continue
        if quantity_value is None or due is None or minutes is None:
            invalid_count += 1
            continue
        quantity = max(1, int(quantity_value))
        code = f"XLSX-{digest[:8].upper()}-{row_number:04d}"
        material_id = material[:90] or f"ROW-{row_number}"
        assigned = person_codes.get(current_person)
        orders.append({
            "externalId": code,
            "code": code,
            "productionOrderId": f"{material_id}-{row_number}"[:128],
            "quantity": quantity,
            "dueAt": due,
            "priority": 50,
            "productRevisionId": f"MATERIAL-{material_id}"[:128],
            "routingRevisionId": "TURNING-ROUTE-INFERRED",
            "bomRevisionId": f"BOM-{material_id}"[:128],
            "drawingRevisionIds": [],
            "status": "RELEASED",
            "version": 1,
            "materialReady": True,
            "qualityHold": False,
            "operations": [{
                "sequence": 10,
                "operationCode": "TURN",
                "operationName": f"车削 · {product_name or material_id}"[:160],
                "workCenterId": "WC-LATHE-01",
                "plannedQuantity": quantity,
                "status": "PENDING",
                "assignedResourceId": assigned,
                "goodQuantity": 0,
                "scrapQuantity": 0,
                "minutesPerUnit": minutes,
                "setupMinutes": 0,
            }],
        })

    now = observed_at or datetime.now(UTC)
    if now.tzinfo is None:
        now = now.replace(tzinfo=UTC)
    snapshot = {
        "sourceSystem": "TURNING_PLAN_WORKBOOK_IMPORT",
        "workshopId": workshop_id.strip(),
        "sourceRevision": f"turning-plan-{digest[:20]}",
        "observedAt": now.isoformat(),
        "workOrders": orders,
        "resources": resources,
    }
    issues: list[dict[str, Any]] = [{
        "severity": "WARNING",
        "sheet": plan_sheet.title,
        "row": None,
        "field": "工单号/工序号",
        "message": "源表没有工单号和工序号；预览按文件哈希与行号生成稳定工单号，并映射为车削工序。",
    }]
    if completed_count:
        issues.append({
            "severity": "WARNING", "sheet": plan_sheet.title, "row": None,
            "field": "状态", "message": f"已排除 {completed_count} 条标记为完成的历史记录。",
        })
    if invalid_count:
        issues.append({
            "severity": "WARNING", "sheet": plan_sheet.title, "row": None,
            "field": "数量/计划完成/单件工时",
            "message": f"已跳过 {invalid_count} 条缺少排产必要字段的记录。",
        })
    if not resources:
        issues.append({
            "severity": "ERROR", "sheet": None, "row": None, "field": "人员能力",
            "message": "没有识别到可用于排产的人员产能。",
        })
    if not orders:
        issues.append({
            "severity": "ERROR", "sheet": plan_sheet.title, "row": None, "field": None,
            "message": "没有识别到未完成且字段完整的计划记录。",
        })
    errors = sum(item["severity"] == "ERROR" for item in issues)
    return {
        "valid": errors == 0,
        "filename": Path(filename).name,
        "previewFingerprint": snapshot_fingerprint(snapshot),
        "stats": {
            "workOrderCount": len(orders),
            "operationCount": len(orders),
            "resourceCount": len(resources),
            "errorCount": errors,
            "warningCount": len(issues) - errors,
            "completedExcludedCount": completed_count,
            "invalidSkippedCount": invalid_count,
        },
        "issues": issues,
        "snapshot": snapshot,
        "formula": "工序需求分钟 = 计划数量 × 单件工时",
        "mappingMode": "TURNING_PLAN_ADAPTER",
    }
