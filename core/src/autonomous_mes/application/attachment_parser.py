"""Bounded, local-only extraction for user-selected manufacturing attachments."""

import csv
import io
import json
import re
import zipfile
from hashlib import sha256
from pathlib import Path
from threading import Lock
from typing import Any

from docx import Document
from openpyxl import load_workbook  # type: ignore[import-untyped]
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader
from rapidocr import RapidOCR

from autonomous_mes.domain.errors import ValidationError

MAX_ATTACHMENT_BYTES = 15 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 30_000
MAX_PDF_PAGES = 100
MAX_SHEETS = 20
MAX_ROWS_PER_SHEET = 1_000
MAX_COLUMNS_PER_SHEET = 256
MAX_ARCHIVE_MEMBERS = 2_000
MAX_ARCHIVE_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_IMAGE_PIXELS = 40_000_000
SUPPORTED_SUFFIXES = frozenset({
    ".txt", ".md", ".log", ".json", ".yaml", ".yml", ".xml",
    ".csv", ".xlsx", ".docx", ".pdf", ".png", ".jpg", ".jpeg",
})
TEXT_SUFFIXES = frozenset({".txt", ".md", ".log", ".json", ".yaml", ".yml", ".xml"})
IMAGE_SUFFIXES = frozenset({".png", ".jpg", ".jpeg"})

_ocr_engine: RapidOCR | None = None
_ocr_lock = Lock()


def _clean_text(value: str) -> str:
    value = value.replace("\x00", "")
    value = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f]", "", value)
    value = re.sub(r"[ \t]+\n", "\n", value)
    value = re.sub(r"\n{4,}", "\n\n\n", value)
    return value.strip()


def _bounded(value: str) -> tuple[str, bool]:
    cleaned = _clean_text(value)
    return cleaned[:MAX_EXTRACTED_CHARACTERS], len(cleaned) > MAX_EXTRACTED_CHARACTERS


def _decode_text(content: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            return content.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise ValidationError("text attachment encoding is not supported")


def _validate_ooxml(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            expanded = sum(member.file_size for member in members)
    except (OSError, zipfile.BadZipFile) as exc:
        raise ValidationError("Office attachment is invalid or damaged") from exc
    if len(members) > MAX_ARCHIVE_MEMBERS or expanded > MAX_ARCHIVE_EXPANDED_BYTES:
        raise ValidationError("Office attachment expands beyond the safe parsing limit")


def _row_text(values: list[str]) -> str:
    while values and not values[-1]:
        values.pop()
    return "\t".join(values)


def _representative_rows(table: Any, limit: int = 2) -> list[str]:
    if not table.data.grid:
        return []
    header = [cell.text.strip() for cell in table.data.grid[0]]
    results: list[str] = []
    for row in table.data.grid[1:]:
        values = [cell.text.strip() for cell in row]
        nonempty = [value for value in values if value]
        if len(nonempty) < 2 or len(set(nonempty)) == 1:
            continue
        pairs = [
            f"{header[index] or f'第{index + 1}列'}={value}"
            for index, value in enumerate(values)
            if value and index < len(header)
        ]
        if pairs:
            results.append("；".join(pairs[:12]))
        if len(results) >= limit:
            break
    return results


def _docling_spreadsheet_text(
    content: bytes, filename: str
) -> tuple[str, dict[str, Any], list[str]]:
    # Imported lazily so non-spreadsheet attachments do not pay Docling's startup cost.
    from docling.backend.msexcel_backend import MsExcelDocumentBackend
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.document import InputDocument

    input_document = InputDocument(
        io.BytesIO(content),
        format=InputFormat.XLSX,
        backend=MsExcelDocumentBackend,
        filename=filename,
    )
    backend = input_document._backend
    if not isinstance(backend, MsExcelDocumentBackend):
        raise TypeError("Docling selected an unexpected spreadsheet backend")
    if not input_document.valid or not backend.is_valid():
        raise ValueError("Docling rejected the workbook")
    document = backend.convert()
    group_names = {group.self_ref: group.name for group in document.groups}
    sheets: list[dict[str, Any]] = []
    tables_by_sheet: dict[str, list[Any]] = {
        group.name: [] for group in document.groups if str(group.label) == "sheet"
    }
    for table in document.tables:
        sheet_name = group_names.get(table.parent.cref) if table.parent is not None else None
        if sheet_name is not None:
            tables_by_sheet.setdefault(sheet_name, []).append(table)

    summary_lines = [f"# 工作簿概览：{len(tables_by_sheet)} 个工作表"]
    for sheet_name, tables in tables_by_sheet.items():
        populated_rows = sum(table.data.num_rows for table in tables)
        max_columns = max((table.data.num_cols for table in tables), default=0)
        first_grid = tables[0].data.grid if tables else []
        fields = [cell.text.strip() for cell in first_grid[0] if cell.text.strip()] \
            if first_grid else []
        profile = {
            "name": sheet_name,
            "regionCount": len(tables),
            "populatedRowCount": populated_rows,
            "maxColumnCount": max_columns,
            "fields": fields[:30],
        }
        sheets.append(profile)
        field_text = "、".join(fields[:16]) or "未识别到稳定表头"
        summary_lines.append(
            f"- {sheet_name}：{len(tables)} 个连续数据区，约 {populated_rows} 行；字段：{field_text}"
        )

    all_fields = {field for sheet in sheets for field in sheet["fields"]}
    missing_mapping_fields = []
    if "工单号" not in all_fields:
        missing_mapping_fields.append("未发现明确标为“工单号”的唯一业务主键")
    if not {"工序号", "工序顺序", "工序序号"}.intersection(all_fields):
        missing_mapping_fields.append("未发现工序顺序字段")
    if not {"工作中心", "资源编号", "设备编号"}.intersection(all_fields):
        missing_mapping_fields.append("未发现工作中心或资源编号字段")
    if missing_mapping_fields:
        summary_lines.append(
            "- MES 映射判断：已理解为制造计划/人员能力数据；"
            + "；".join(missing_mapping_fields)
            + "，需要字段映射后才能写入。"
        )

    # Samples from every sheet are placed before longer excerpts. Field=value prose is
    # materially more reliable than raw TSV for small local language models.
    for sheet_name, tables in tables_by_sheet.items():
        summary_lines.append(f"\n## {sheet_name} 样例")
        samples = _representative_rows(tables[0], 2) if tables else []
        summary_lines.extend(f"- 记录：{sample}" for sample in samples)
        if not samples:
            summary_lines.append("- 未找到可稳定映射的样例记录。")
    summary = "\n".join(summary_lines)

    remaining = max(0, MAX_EXTRACTED_CHARACTERS - len(summary) - 2)
    per_sheet_budget = remaining // max(1, len(tables_by_sheet))
    excerpts: list[str] = []
    for sheet_name, tables in tables_by_sheet.items():
        lines = [f"## 工作表：{sheet_name}"]
        used = len(lines[0])
        for region_index, table in enumerate(tables, start=1):
            heading = f"### 连续数据区 {region_index}（{table.data.num_rows} 行 × {table.data.num_cols} 列）"
            if used + len(heading) + 1 > per_sheet_budget:
                break
            lines.append(heading)
            used += len(heading) + 1
            for row in table.data.grid:
                line = _row_text([cell.text for cell in row])
                if not line:
                    continue
                if used + len(line) + 1 > per_sheet_budget:
                    break
                lines.append(line)
                used += len(line) + 1
            if used >= per_sheet_budget:
                break
        excerpts.append("\n".join(lines))

    text = f"{summary}\n\n" + "\n\n".join(excerpts)
    metadata = {
        "engine": "Docling",
        "engineVersion": "2.126.0",
        "sheetCount": len(tables_by_sheet),
        "regionCount": len(document.tables),
        "rowCount": sum(item["populatedRowCount"] for item in sheets),
        "sheets": sheets,
        "mappingStatus": "REQUIRES_FIELD_MAPPING" if missing_mapping_fields else "MAPPABLE",
        "mappingGaps": missing_mapping_fields,
        "summary": summary,
    }
    return text, metadata, []


def _openpyxl_spreadsheet_text(
    content: bytes,
) -> tuple[str, dict[str, Any], list[str]]:
    _validate_ooxml(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError("XLSX attachment is invalid or damaged") from exc
    chunks: list[str] = []
    row_count = 0
    sheet_names = workbook.sheetnames[:MAX_SHEETS]
    for sheet_name in sheet_names:
        sheet = workbook[sheet_name]
        chunks.append(f"## Sheet: {sheet_name}")
        max_column = min(sheet.max_column, MAX_COLUMNS_PER_SHEET)
        for row_index, row in enumerate(
            sheet.iter_rows(max_col=max_column, values_only=True)
        ):
            if row_index >= MAX_ROWS_PER_SHEET:
                break
            line = _row_text(["" if value is None else str(value) for value in row])
            if line:
                chunks.append(line)
                row_count += 1
    warnings = []
    if len(workbook.sheetnames) > MAX_SHEETS:
        warnings.append("工作表数量超过解析上限，仅提取前 20 个。")
    return "\n".join(chunks), {"sheetCount": len(workbook.sheetnames), "rowCount": row_count}, warnings


def _spreadsheet_text(
    content: bytes, filename: str
) -> tuple[str, dict[str, Any], list[str]]:
    _validate_ooxml(content)
    try:
        return _docling_spreadsheet_text(content, filename)
    except Exception:  # noqa: BLE001 - bounded openpyxl extraction remains available
        text, metadata, warnings = _openpyxl_spreadsheet_text(content)
        warnings.insert(0, "结构化工作表识别未完成，已切换为基础单元格提取。")
        metadata["engine"] = "openpyxl-fallback"
        return text, metadata, warnings


def _csv_text(content: bytes) -> tuple[str, dict[str, Any], list[str]]:
    decoded, encoding = _decode_text(content)
    rows = list(csv.reader(io.StringIO(decoded)))
    bounded_rows = rows[:MAX_ROWS_PER_SHEET]
    text = "\n".join("\t".join(cell for cell in row) for row in bounded_rows)
    warnings = ["CSV 行数超过解析上限，仅提取前 1000 行。"] if len(rows) > len(bounded_rows) else []
    return text, {"rowCount": len(rows), "encoding": encoding}, warnings


def _document_text(content: bytes) -> tuple[str, dict[str, Any], list[str]]:
    _validate_ooxml(content)
    try:
        document = Document(io.BytesIO(content))
    except Exception as exc:
        raise ValidationError("DOCX attachment is invalid or damaged") from exc
    chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table_index, table in enumerate(document.tables, start=1):
        chunks.append(f"## Table {table_index}")
        for row in table.rows:
            chunks.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(chunks), {
        "paragraphCount": len(document.paragraphs), "tableCount": len(document.tables),
    }, []


def _ocr(content: bytes) -> tuple[str, dict[str, Any]]:
    global _ocr_engine
    try:
        with Image.open(io.BytesIO(content)) as image:
            width, height = image.size
            if width * height > MAX_IMAGE_PIXELS:
                raise ValidationError("image dimensions exceed the safe OCR limit")
            image.verify()
    except (OSError, UnidentifiedImageError) as exc:
        raise ValidationError("image attachment could not be decoded") from exc
    with _ocr_lock:
        if _ocr_engine is None:
            _ocr_engine = RapidOCR()
        result: Any = _ocr_engine(content)
    texts = list(getattr(result, "txts", ()) or ())
    scores = [float(score) for score in (getattr(result, "scores", ()) or ())]
    confidence = round(sum(scores) / len(scores), 4) if scores else None
    return "\n".join(texts), {"ocrLineCount": len(texts), "ocrMeanConfidence": confidence}


def _pdf_text(content: bytes) -> tuple[str, dict[str, Any], list[str]]:
    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as exc:
        raise ValidationError("PDF attachment is invalid, encrypted, or damaged") from exc
    chunks: list[str] = []
    ocr_pages = 0
    pages = list(reader.pages[:MAX_PDF_PAGES])
    for page_index, page in enumerate(pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            image_texts: list[str] = []
            try:
                for image in list(page.images)[:4]:
                    extracted, _ = _ocr(image.data)
                    if extracted:
                        image_texts.append(extracted)
            except Exception:  # noqa: BLE001 - malformed embedded images are optional OCR input
                image_texts = []
            text = "\n".join(image_texts)
            if text:
                ocr_pages += 1
        if text:
            chunks.append(f"## Page {page_index}\n{text}")
    warnings = []
    if len(reader.pages) > MAX_PDF_PAGES:
        warnings.append("PDF 页数超过解析上限，仅提取前 100 页。")
    if not chunks:
        warnings.append("PDF 中没有提取到可读文字。")
    return "\n".join(chunks), {"pageCount": len(reader.pages), "ocrPageCount": ocr_pages}, warnings


def parse_attachment(content: bytes, filename: str) -> dict[str, Any]:
    safe_name = Path(filename).name
    suffix = Path(safe_name).suffix.lower()
    if not safe_name or suffix not in SUPPORTED_SUFFIXES:
        raise ValidationError("attachment type is not supported")
    if not content:
        raise ValidationError("attachment is empty")
    if len(content) > MAX_ATTACHMENT_BYTES:
        raise ValidationError("attachment exceeds 15 MB limit")

    warnings: list[str] = []
    metadata: dict[str, Any] = {}
    parser = "LOCAL_TEXT"
    kind = "TEXT"
    if suffix in TEXT_SUFFIXES:
        text, encoding = _decode_text(content)
        metadata = {"encoding": encoding}
        if suffix == ".json":
            try:
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
            except json.JSONDecodeError:
                warnings.append("JSON 语法无效，已按普通文字提取。")
    elif suffix == ".csv":
        kind, parser = "TABLE", "LOCAL_CSV"
        text, metadata, warnings = _csv_text(content)
    elif suffix == ".xlsx":
        kind = "WORKBOOK"
        text, metadata, warnings = _spreadsheet_text(content, safe_name)
        parser = (
            "LOCAL_DOCLING_XLSX"
            if metadata.get("engine") == "Docling"
            else "LOCAL_OPENPYXL"
        )
    elif suffix == ".docx":
        kind, parser = "DOCUMENT", "LOCAL_PYTHON_DOCX"
        text, metadata, warnings = _document_text(content)
    elif suffix == ".pdf":
        kind, parser = "DOCUMENT", "LOCAL_PYPDF_OCR"
        text, metadata, warnings = _pdf_text(content)
    elif suffix in IMAGE_SUFFIXES:
        kind, parser = "IMAGE_OCR", "LOCAL_RAPIDOCR"
        try:
            text, metadata = _ocr(content)
        except Exception as exc:
            raise ValidationError("image attachment could not be decoded") from exc
        if not text:
            warnings.append("图片中没有识别到文字；当前版本不判断外观、缺陷或图纸几何。")
    else:  # pragma: no cover - guarded by SUPPORTED_SUFFIXES
        raise ValidationError("attachment type is not supported")

    text, truncated = _bounded(text)
    if truncated:
        warnings.append("提取内容超过 30000 字符，已截断。")
    return {
        "name": safe_name,
        "kind": kind,
        "parser": parser,
        "sha256": sha256(content).hexdigest(),
        "status": "PARSED" if text else "NO_TEXT",
        "text": text,
        "characterCount": len(text),
        "truncated": truncated,
        "metadata": metadata,
        "summary": str(metadata.get("summary", "")),
        "warnings": warnings,
        "localOnly": True,
    }
