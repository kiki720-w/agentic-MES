from __future__ import annotations

import io
import json
import re
from hashlib import sha256
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook  # type: ignore[import-untyped]

from autonomous_mes.domain.errors import InvalidTransition, ValidationError

from .scheduling import (
    RegisterPlanningResourceCommand,
    SchedulingApplicationService,
    UpdatePlanningResourceCommand,
)

MAX_AGENT_WORKBOOK_BYTES = 15 * 1024 * 1024
MAX_SCAN_ROWS = 5_000
MAX_SCAN_COLUMNS = 64

_PERSON_HEADERS = {"人员", "姓名", "员工", "员工姓名", "操作员", "人员姓名"}
_CAPABILITY_HEADERS = {"加工种类", "能力", "技能", "工种", "工序", "可加工产品"}
_DAILY_HEADERS = {"8小时工时", "日工时", "正常工时", "日能力", "日产能", "标准产能"}
_OVERTIME_HEADERS = {"加班3小时工时", "加班工时", "加班产能", "最大产能"}


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _header(value: Any) -> str:
    return re.sub(r"[\s：:()（）/\\_-]+", "", _text(value))


def _number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value) if float(value) > 0 else None
    matches = re.findall(r"\d+(?:\.\d+)?", _text(value).replace(",", ""))
    if not matches:
        return None
    parsed = float(matches[-1])
    return parsed if parsed > 0 else None


def _stable_person_code(name: str) -> str:
    normalized = re.sub(r"\s+", "", name).casefold()
    return f"PERSON-{sha256(normalized.encode('utf-8')).hexdigest()[:12].upper()}"


def capacity_plan_fingerprint(plan: dict[str, Any]) -> str:
    canonical = {
        key: value for key, value in plan.items() if key != "previewFingerprint"
    }
    encoded = json.dumps(
        canonical, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return sha256(encoded).hexdigest()


def _find_capacity_table(workbook: Any) -> tuple[Any, int, dict[str, int]] | None:
    for sheet in workbook.worksheets[:20]:
        for row_number, row in enumerate(
            sheet.iter_rows(
                min_row=1,
                max_row=min(sheet.max_row, 30),
                max_col=min(sheet.max_column, MAX_SCAN_COLUMNS),
                values_only=True,
            ),
            start=1,
        ):
            indexes: dict[str, int] = {}
            for index, value in enumerate(row):
                normalized = _header(value)
                if normalized in _PERSON_HEADERS:
                    indexes["person"] = index
                elif normalized in _CAPABILITY_HEADERS:
                    indexes["capability"] = index
                elif normalized in _DAILY_HEADERS:
                    indexes["daily"] = index
                elif normalized in _OVERTIME_HEADERS:
                    indexes["overtime"] = index
            if {"person", "capability", "daily"}.issubset(indexes):
                return sheet, row_number, indexes
    return None


def preview_capacity_workbook(
    content: bytes,
    filename: str,
    workshop_id: str,
    existing_resources: list[dict[str, Any]],
    work_center_id: str = "WC-LATHE-01",
) -> dict[str, Any]:
    if not filename.lower().endswith(".xlsx"):
        return _invalid_preview(filename, content, "当前人员能力导入支持 .xlsx 文件。")
    if not content or len(content) > MAX_AGENT_WORKBOOK_BYTES:
        return _invalid_preview(filename, content, "文件为空或超过 15 MB。")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError("XLSX attachment is invalid or damaged") from exc

    located = _find_capacity_table(workbook)
    if located is None:
        return _invalid_preview(
            filename,
            content,
            "未找到同时包含人员、加工种类和日工时/日产能的能力表。",
        )
    sheet, header_row, indexes = located
    people: dict[str, dict[str, Any]] = {}
    current_person = ""
    for row in sheet.iter_rows(
        min_row=header_row + 1,
        max_row=min(sheet.max_row, header_row + MAX_SCAN_ROWS),
        max_col=min(sheet.max_column, MAX_SCAN_COLUMNS),
        values_only=True,
    ):
        person = _text(row[indexes["person"]])
        if person:
            current_person = person
        if not current_person:
            continue
        capability = _text(row[indexes["capability"]])
        daily = _number(row[indexes["daily"]])
        overtime = _number(row[indexes["overtime"]]) if "overtime" in indexes else None
        if not capability and daily is None and overtime is None:
            continue
        record = people.setdefault(
            current_person,
            {"capabilities": {"TURN"}, "daily": [], "overtime": []},
        )
        if capability:
            record["capabilities"].add(capability)
        if daily is not None:
            record["daily"].append(daily)
        if overtime is not None:
            record["overtime"].append(overtime)

    issues: list[dict[str, Any]] = []
    by_code = {str(item["code"]).casefold(): item for item in existing_resources}
    by_name: dict[str, list[dict[str, Any]]] = {}
    for item in existing_resources:
        by_name.setdefault(str(item["name"]).strip().casefold(), []).append(item)
    actions: list[dict[str, Any]] = []
    for name, record in people.items():
        if not record["daily"]:
            issues.append({
                "severity": "WARNING",
                "sheet": sheet.title,
                "field": "日工时/日产能",
                "message": f"{name} 没有可用的正常产能数值，本次跳过，可在产能管理中人工补录。",
            })
            continue
        code = _stable_person_code(name)
        daily_capacity = max(record["daily"])
        overtime_capacity = max(record["overtime"] or [daily_capacity])
        overtime_capacity = max(daily_capacity, overtime_capacity)
        same_name = by_name.get(name.casefold(), [])
        existing = by_code.get(code.casefold())
        if existing is None and len(same_name) == 1:
            existing = same_name[0]
            code = str(existing["code"])
        elif existing is None and len(same_name) > 1:
            issues.append({
                "severity": "ERROR",
                "sheet": sheet.title,
                "field": "人员",
                "message": f"现有产能管理中有多个同名人员“{name}”，无法自动确定更新对象。",
            })
            continue
        capabilities = sorted(record["capabilities"])
        target = {
            "code": code,
            "name": name,
            "resourceType": "PERSON",
            "workshopId": workshop_id,
            "workCenterId": work_center_id,
            "dailyCapacityMinutes": daily_capacity,
            "overtimeCapacityMinutes": overtime_capacity,
            "capabilityCodes": capabilities,
            "active": True,
        }
        if existing is None:
            action_type = "CREATE"
        else:
            comparable = {
                key: existing.get(key)
                for key in (
                    "code", "name", "resourceType", "workshopId", "workCenterId",
                    "dailyCapacityMinutes", "overtimeCapacityMinutes", "capabilityCodes", "active",
                )
            }
            comparable["capabilityCodes"] = sorted(comparable.get("capabilityCodes") or [])
            action_type = "UNCHANGED" if comparable == target else "UPDATE"
            target["resourceId"] = existing["resourceId"]
            target["expectedVersion"] = existing["version"]
        actions.append({"action": action_type, "resource": target})

    stats = {
        "personCount": len(actions),
        "createCount": sum(item["action"] == "CREATE" for item in actions),
        "updateCount": sum(item["action"] == "UPDATE" for item in actions),
        "unchangedCount": sum(item["action"] == "UNCHANGED" for item in actions),
        "skippedCount": sum(item["severity"] == "WARNING" for item in issues),
        "errorCount": sum(item["severity"] == "ERROR" for item in issues),
    }
    plan: dict[str, Any] = {
        "valid": bool(actions) and stats["errorCount"] == 0,
        "intent": "UPSERT_PLANNING_RESOURCES",
        "source": {
            "filename": filename,
            "sha256": sha256(content).hexdigest(),
            "sheet": sheet.title,
            "headerRow": header_row,
        },
        "workshopId": workshop_id,
        "workCenterId": work_center_id,
        "stats": stats,
        "issues": issues,
        "assumptions": [
            "未提供工号时，系统按姓名生成稳定人员编号；再次导入会更新同一人员。",
            "同一人员有多项加工能力时合并能力，正常与加班产能分别取非空最大值。",
            "车工名单中的人员统一增加 TURN 工序能力，表内加工种类继续作为细分能力保留。",
            f"该车工能力表映射到工作中心 {work_center_id}，可在产能管理中人工修改。",
        ],
        "actions": actions,
    }
    plan["previewFingerprint"] = capacity_plan_fingerprint(plan)
    return plan


def _invalid_preview(filename: str, content: bytes, message: str) -> dict[str, Any]:
    plan: dict[str, Any] = {
        "valid": False,
        "intent": "UNRECOGNIZED_WORKBOOK",
        "source": {"filename": filename, "sha256": sha256(content).hexdigest()},
        "stats": {
            "personCount": 0, "createCount": 0, "updateCount": 0,
            "unchangedCount": 0, "skippedCount": 0, "errorCount": 1,
        },
        "issues": [{"severity": "ERROR", "sheet": None, "field": None, "message": message}],
        "assumptions": [],
        "actions": [],
    }
    plan["previewFingerprint"] = capacity_plan_fingerprint(plan)
    return plan


def execute_capacity_plan(
    plan: dict[str, Any],
    preview_fingerprint: str,
    service: SchedulingApplicationService,
    actor_id: str,
) -> dict[str, Any]:
    if not plan.get("valid") or plan.get("intent") != "UPSERT_PLANNING_RESOURCES":
        raise ValidationError("capacity import plan is not valid")
    if capacity_plan_fingerprint(plan) != preview_fingerprint:
        raise ValidationError("preview fingerprint changed; preview the file again")
    workshop_id = str(plan.get("workshopId", "")).strip()
    current = service.list_resources(workshop_id)
    current_by_id = {str(item["resourceId"]): item for item in current}
    current_by_code = {str(item["code"]).casefold(): item for item in current}

    for action in plan.get("actions", []):
        kind = str(action.get("action", ""))
        resource = action.get("resource") or {}
        if kind == "CREATE" and str(resource.get("code", "")).casefold() in current_by_code:
            raise InvalidTransition("capacity data changed after preview; preview the file again")
        if kind in {"UPDATE", "UNCHANGED"}:
            existing = current_by_id.get(str(resource.get("resourceId", "")))
            if existing is None or existing["version"] != resource.get("expectedVersion"):
                raise InvalidTransition("capacity data changed after preview; preview the file again")

    results: list[dict[str, Any]] = []
    for action in plan.get("actions", []):
        kind = str(action["action"])
        resource = action["resource"]
        if kind == "CREATE":
            saved = service.register_resource(RegisterPlanningResourceCommand(
                resource["code"], resource["name"], resource["resourceType"],
                resource["workshopId"], resource["workCenterId"],
                resource["dailyCapacityMinutes"], resource["overtimeCapacityMinutes"],
                list(resource["capabilityCodes"]), actor_id, str(uuid4()),
            ))
        elif kind == "UPDATE":
            saved = service.update_resource(UpdatePlanningResourceCommand(
                resource["resourceId"], resource["expectedVersion"], resource["name"],
                resource["workCenterId"], resource["dailyCapacityMinutes"],
                resource["overtimeCapacityMinutes"], list(resource["capabilityCodes"]),
                bool(resource["active"]), actor_id, str(uuid4()),
            ))
        elif kind == "UNCHANGED":
            saved = current_by_id[str(resource["resourceId"])]
        else:
            raise ValidationError("capacity import plan contains an unsupported action")
        results.append({"action": kind, "resource": saved})

    verified_resources = service.list_resources(workshop_id)
    verified_by_code = {str(item["code"]).casefold(): item for item in verified_resources}
    expected_codes: list[str] = []
    for action in plan["actions"]:
        expected = action["resource"]
        code = str(expected["code"]).casefold()
        expected_codes.append(code)
        actual = verified_by_code.get(code)
        if actual is None or any((
            actual["name"] != expected["name"],
            actual["resourceType"] != expected["resourceType"],
            actual["workshopId"] != expected["workshopId"],
            actual["workCenterId"] != expected["workCenterId"],
            actual["dailyCapacityMinutes"] != expected["dailyCapacityMinutes"],
            actual["overtimeCapacityMinutes"] != expected["overtimeCapacityMinutes"],
            sorted(actual["capabilityCodes"]) != sorted(expected["capabilityCodes"]),
            actual["active"] != expected["active"],
        )):
            raise InvalidTransition("capacity import write-back verification failed")
    return {
        "status": "EXECUTED_AND_VERIFIED",
        "workshopId": workshop_id,
        "source": plan["source"],
        "stats": plan["stats"],
        "results": results,
        "verifiedCount": len(expected_codes),
    }
