import csv
import io
import json
from datetime import UTC, date, datetime, time
from hashlib import sha256
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS_PER_SHEET = 5_000

SHEET_ALIASES = {
    "orders": {"工单", "工单表", "orders", "order"},
    "operations": {"工序", "工艺路线", "operations", "routing"},
    "resources": {"产能", "资源", "人员与工作单元", "resources", "capacity"},
}

FIELD_ALIASES = {
    "order_code": {"工单号", "工单编号", "订单号", "workorder", "workordercode", "code"},
    "production_order": {"生产订单", "生产订单号", "productionorder", "productionorderid"},
    "quantity": {"数量", "计划数量", "需求数量", "quantity", "qty"},
    "due_at": {"交期", "交货期", "要求完成时间", "dueat", "duedate"},
    "priority": {"优先级", "priority"},
    "product": {"产品", "产品编码", "工件", "零件", "product", "productrevisionid"},
    "routing": {"工艺版本", "工艺路线版本", "routingrevisionid"},
    "bom": {"bom版本", "bomrevisionid"},
    "status": {"状态", "status"},
    "material_ready": {"物料齐套", "materialready"},
    "quality_hold": {"质量冻结", "qualityhold"},
    "sequence": {"工序号", "工序顺序", "序号", "sequence"},
    "operation_code": {"工序编码", "工序代码", "operationcode"},
    "operation_name": {"工序名称", "operationname"},
    "work_center": {"工作中心", "工作中心编码", "workcenter", "workcenterid"},
    "minutes_per_unit": {"单件工时", "单件分钟", "标准工时", "minutesperunit", "unitminutes"},
    "setup_minutes": {"准备工时", "换型时间", "setupminutes"},
    "planned_quantity": {"工序数量", "计划数量", "plannedquantity"},
    "resource_code": {"资源编码", "人员编码", "设备编码", "resourcecode", "code"},
    "resource_name": {"资源名称", "姓名", "设备名称", "resourcename", "name"},
    "resource_type": {"资源类型", "类型", "resourcetype"},
    "daily_capacity": {"日常产能分钟", "正常产能分钟", "日产能", "dailycapacityminutes"},
    "overtime_capacity": {"含加班产能分钟", "最大产能分钟", "overtimecapacityminutes"},
    "capabilities": {"能力编码", "技能", "可做工序", "capabilitycodes", "skills"},
    "active": {"启用", "有效", "active"},
    "state": {"设备状态", "state"},
}


def preview_spreadsheet(
    content: bytes, filename: str, workshop_id: str, observed_at: datetime | None = None
) -> dict[str, Any]:
    if not content:
        return _failed("文件为空")
    if len(content) > MAX_FILE_BYTES:
        return _failed("文件超过 5 MB 安全限制")
    suffix = Path(filename).suffix.lower()
    issues: list[dict[str, Any]] = []
    try:
        if suffix == ".xlsx":
            tables = _xlsx_tables(content, issues)
        elif suffix == ".csv":
            tables = _csv_tables(content, issues)
        else:
            return _failed("目前业务导入支持 .xlsx 和 .csv；图片/PDF 可作为对话附件预览")
    except Exception as exc:  # noqa: BLE001 - parser boundary returns a safe validation issue
        return _failed(f"无法解析文件：{exc}")

    orders = _parse_orders(tables.get("orders", []), issues)
    operations = _parse_operations(tables.get("operations", []), issues)
    resources = _parse_resources(tables.get("resources", []), issues)
    by_code = {item["code"]: item for item in orders}
    for source_row, operation in operations:
        order_code = operation.pop("_orderCode")
        order = by_code.get(order_code)
        if order is None:
            _issue(issues, "ERROR", "工序", source_row, "工单号", f"找不到工单 {order_code}")
            continue
        if operation["plannedQuantity"] == 0:
            operation["plannedQuantity"] = order["quantity"]
        order["operations"].append(operation)
    for order in orders:
        order["operations"].sort(key=lambda item: item["sequence"])
        if not order["operations"]:
            _issue(issues, "ERROR", "工单", None, "工序", f"工单 {order['code']} 没有工序")
    if not orders:
        _issue(issues, "ERROR", "工单", None, None, "没有可导入的工单")
    if not resources:
        _issue(issues, "ERROR", "产能", None, None, "没有可用于排产的产能资源")

    now = observed_at or datetime.now(UTC)
    if now.tzinfo is None:
        now = now.replace(tzinfo=UTC)
    file_digest = sha256(content).hexdigest()
    snapshot = {
        "sourceSystem": "HUMAN_SPREADSHEET_IMPORT",
        "workshopId": workshop_id.strip(),
        "sourceRevision": f"spreadsheet-{file_digest[:20]}",
        "observedAt": now.isoformat(),
        "workOrders": orders,
        "resources": resources,
    }
    encoded = json.dumps(snapshot, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    fingerprint = sha256(encoded.encode()).hexdigest()
    errors = sum(item["severity"] == "ERROR" for item in issues)
    warnings = sum(item["severity"] == "WARNING" for item in issues)
    return {
        "valid": errors == 0,
        "filename": Path(filename).name,
        "previewFingerprint": fingerprint,
        "stats": {
            "workOrderCount": len(orders),
            "operationCount": sum(len(item["operations"]) for item in orders),
            "resourceCount": len(resources),
            "errorCount": errors,
            "warningCount": warnings,
        },
        "issues": issues[:200],
        "snapshot": snapshot,
        "formula": "工序需求分钟 = 准备工时 + (计划数量 - 良品 - 报废) × 单件工时",
    }


def snapshot_fingerprint(snapshot: dict[str, Any]) -> str:
    encoded = json.dumps(snapshot, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return sha256(encoded.encode()).hexdigest()


def build_spreadsheet_template() -> bytes:
    workbook = Workbook()
    orders = workbook.active
    orders.title = "工单"
    orders.append(
        ["工单号", "生产订单号", "产品编码", "数量", "交期", "优先级", "物料齐套", "质量冻结"]
    )
    orders.append(
        [
            "WO-DEMO-001",
            "PO-DEMO-001",
            "PART-A-REV1",
            100,
            "2026-09-20T17:00:00+10:00",
            80,
            "是",
            "否",
        ]
    )
    operations = workbook.create_sheet("工序")
    operations.append(
        ["工单号", "工序号", "工序编码", "工序名称", "工作中心", "单件工时", "准备工时", "工序数量"]
    )
    operations.append(["WO-DEMO-001", 10, "TURN", "车削", "WC-LATHE-01", 12, 30, 100])
    operations.append(["WO-DEMO-001", 20, "GRIND", "磨削", "WC-GRIND-01", 8, 20, 100])
    resources = workbook.create_sheet("产能")
    resources.append(
        [
            "资源编码",
            "资源名称",
            "资源类型",
            "工作中心",
            "日常产能分钟",
            "含加班产能分钟",
            "能力编码",
            "启用",
            "设备状态",
        ]
    )
    resources.append(
        ["OP-001", "车削操作员甲", "人员", "WC-LATHE-01", 480, 600, "TURN", "是", "UNKNOWN"]
    )
    resources.append(
        [
            "CELL-GRIND-01",
            "磨削单元一",
            "工作单元",
            "WC-GRIND-01",
            480,
            600,
            "GRIND",
            "是",
            "UNKNOWN",
        ]
    )
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="176B74")
        for column in sheet.columns:
            width = min(34, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _xlsx_tables(content: bytes, issues: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    tables: dict[str, list[dict[str, Any]]] = {}
    for kind, aliases in SHEET_ALIASES.items():
        sheet = next(
            (
                item
                for item in workbook.worksheets
                if _norm(item.title) in {_norm(a) for a in aliases}
            ),
            None,
        )
        if sheet is None:
            _issue(issues, "ERROR", kind, None, None, f"缺少工作表：{min(aliases)}")
            continue
        rows = list(sheet.iter_rows(values_only=True, max_row=MAX_ROWS_PER_SHEET + 1))
        if len(rows) > MAX_ROWS_PER_SHEET:
            _issue(issues, "ERROR", sheet.title, None, None, "单表超过 5000 行限制")
            rows = rows[:MAX_ROWS_PER_SHEET]
        tables[kind] = _rows_to_dicts(rows)
    return tables


def _csv_tables(content: bytes, issues: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    text = content.decode("utf-8-sig")
    rows = list(csv.DictReader(io.StringIO(text)))[:MAX_ROWS_PER_SHEET]
    headers = {_norm(item) for item in (rows[0].keys() if rows else [])}
    if any(_norm(a) in headers for a in FIELD_ALIASES["resource_code"]):
        return {"orders": [], "operations": [], "resources": rows}
    _issue(
        issues,
        "ERROR",
        "CSV",
        None,
        None,
        "CSV 目前仅支持产能资源表；完整排产输入请使用三工作表 XLSX",
    )
    return {"orders": rows, "operations": [], "resources": []}


def _rows_to_dicts(rows: list[tuple[Any, ...]]) -> list[dict[str, Any]]:
    if not rows:
        return []
    headers = [str(item).strip() if item is not None else "" for item in rows[0]]
    return [
        {
            headers[index]: value
            for index, value in enumerate(row)
            if index < len(headers) and headers[index]
        }
        for row in rows[1:]
        if any(value not in (None, "") for value in row)
    ]


def _parse_orders(rows: list[dict[str, Any]], issues: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result = []
    for index, row in enumerate(rows, 2):
        try:
            code = _required(row, "order_code")
            quantity = int(float(_required(row, "quantity")))
            due_at = _datetime_value(_required(row, "due_at"))
            product = str(_value(row, "product", code))
            result.append(
                {
                    "externalId": code,
                    "code": code,
                    "productionOrderId": str(_value(row, "production_order", code)),
                    "quantity": quantity,
                    "dueAt": due_at,
                    "priority": int(float(_value(row, "priority", 50))),
                    "productRevisionId": product,
                    "routingRevisionId": str(_value(row, "routing", f"ROUTE-{product}")),
                    "bomRevisionId": str(_value(row, "bom", f"BOM-{product}")),
                    "drawingRevisionIds": [],
                    "status": str(_value(row, "status", "RELEASED")).upper(),
                    "version": 1,
                    "materialReady": _bool_value(_value(row, "material_ready", True)),
                    "qualityHold": _bool_value(_value(row, "quality_hold", False)),
                    "operations": [],
                }
            )
        except (KeyError, TypeError, ValueError) as exc:
            _issue(issues, "ERROR", "工单", index, None, str(exc))
    return result


def _parse_operations(
    rows: list[dict[str, Any]], issues: list[dict[str, Any]]
) -> list[tuple[int, dict[str, Any]]]:
    result = []
    for index, row in enumerate(rows, 2):
        try:
            planned = _value(row, "planned_quantity", None)
            result.append(
                (
                    index,
                    {
                        "_orderCode": _required(row, "order_code"),
                        "sequence": int(float(_required(row, "sequence"))),
                        "operationCode": _required(row, "operation_code"),
                        "operationName": str(
                            _value(row, "operation_name", _required(row, "operation_code"))
                        ),
                        "workCenterId": _required(row, "work_center"),
                        "plannedQuantity": int(float(planned)) if planned not in (None, "") else 0,
                        "status": str(_value(row, "status", "PENDING")).upper(),
                        "assignedResourceId": None,
                        "goodQuantity": 0,
                        "scrapQuantity": 0,
                        "minutesPerUnit": float(_required(row, "minutes_per_unit")),
                        "setupMinutes": float(_value(row, "setup_minutes", 0)),
                    },
                )
            )
        except (KeyError, TypeError, ValueError) as exc:
            _issue(issues, "ERROR", "工序", index, None, str(exc))
    return result


def _parse_resources(
    rows: list[dict[str, Any]], issues: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    result = []
    type_map = {
        "人员": "PERSON",
        "人": "PERSON",
        "工作单元": "CELL",
        "单元": "CELL",
        "设备": "EQUIPMENT",
    }
    for index, row in enumerate(rows, 2):
        try:
            code = _required(row, "resource_code")
            kind_raw = str(_value(row, "resource_type", "PERSON")).strip()
            kind = type_map.get(kind_raw, kind_raw.upper())
            daily = float(_required(row, "daily_capacity"))
            overtime = float(_value(row, "overtime_capacity", daily))
            capabilities = [
                item.strip().upper()
                for item in str(_value(row, "capabilities", "*")).replace("，", ",").split(",")
                if item.strip()
            ]
            result.append(
                {
                    "externalId": code,
                    "code": code,
                    "name": str(_value(row, "resource_name", code)),
                    "resourceType": kind,
                    "workCenterId": _required(row, "work_center"),
                    "dailyCapacityMinutes": daily,
                    "overtimeCapacityMinutes": overtime,
                    "capabilityCodes": capabilities,
                    "active": _bool_value(_value(row, "active", True)),
                    "state": str(
                        _value(row, "state", "IDLE" if kind == "EQUIPMENT" else "UNKNOWN")
                    ).upper(),
                }
            )
        except (KeyError, TypeError, ValueError) as exc:
            _issue(issues, "ERROR", "产能", index, None, str(exc))
    return result


def _value(row: dict[str, Any], field: str, default: Any = None) -> Any:
    normalized = {_norm(key): value for key, value in row.items()}
    for alias in FIELD_ALIASES[field]:
        value = normalized.get(_norm(alias))
        if value not in (None, ""):
            return value
    return default


def _required(row: dict[str, Any], field: str) -> str:
    value = _value(row, field)
    if value in (None, ""):
        raise ValueError(f"缺少字段：{min(FIELD_ALIASES[field])}")
    return str(value).strip()


def _datetime_value(value: Any) -> str:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time(17, 0))
    else:
        parsed = datetime.fromisoformat(str(value))
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=UTC)
    return parsed.isoformat()


def _bool_value(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() not in {"0", "false", "否", "n", "no", "停用"}


def _norm(value: Any) -> str:
    return "".join(character for character in str(value).strip().lower() if character.isalnum())


def _issue(
    issues: list[dict[str, Any]],
    severity: str,
    sheet: str,
    row: int | None,
    field: str | None,
    message: str,
) -> None:
    issues.append(
        {"severity": severity, "sheet": sheet, "row": row, "field": field, "message": message}
    )


def _failed(message: str) -> dict[str, Any]:
    return {
        "valid": False,
        "filename": None,
        "previewFingerprint": None,
        "stats": {
            "workOrderCount": 0,
            "operationCount": 0,
            "resourceCount": 0,
            "errorCount": 1,
            "warningCount": 0,
        },
        "issues": [
            {"severity": "ERROR", "sheet": "文件", "row": None, "field": None, "message": message}
        ],
        "snapshot": None,
    }
