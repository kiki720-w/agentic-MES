from pathlib import Path
from datetime import date
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\33929\Documents\ChatGPT\mes")
OUT = ROOT / "自主MES产品规划与流程架构.docx"
ASSET = ROOT / "doc_assets"
ASSET.mkdir(exist_ok=True)

NAVY = "17365D"
BLUE = "2E74B5"
CYAN = "2F8F9D"
PALE_BLUE = "EAF2F8"
PALE_CYAN = "E9F5F5"
PALE_GOLD = "FFF4D6"
GOLD = "C58A00"
GRAY = "5B6573"
LIGHT = "F2F4F7"
MID = "D7DEE8"
GREEN = "287D5A"
RED = "A33A3A"
WHITE = "FFFFFF"

FONT_CN = r"C:\Windows\Fonts\msyh.ttc"
FONT_CN_BOLD = r"C:\Windows\Fonts\msyhbd.ttc"


def rgb(hexstr):
    return RGBColor.from_string(hexstr)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa, indent=120):
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_run_font(run, size=11, bold=False, color="000000", name="Microsoft YaHei"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = rgb(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, 9, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, 9, color=GRAY)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run_font(r, 10.5)
    return p


def add_number(doc, text, style="List Number"):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run_font(r, 10.5)
    return p


def add_para(doc, text, bold_lead=None, color="000000", align=None, after=6, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, 10.5, True, color)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, 10.5, False, color)
        r2.italic = italic
    else:
        r = p.add_run(text)
        set_run_font(r, 10.5, False, color)
        r.italic = italic
    return p


def add_callout(doc, title, body, fill=PALE_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, 11, True, accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(body)
    set_run_font(r2, 10.5, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths, aligns=None, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        set_run_font(r, 9.5, True, WHITE)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2:
                set_cell_shading(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            if aligns:
                p.alignment = aligns[i]
            r = p.add_run(str(val))
            set_run_font(r, font_size, color="202733")
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def rounded(draw, xy, fill, outline, radius=20, width=3):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def fit_text(draw, text, box, font_path=FONT_CN, max_size=30, min_size=16, spacing=6):
    x1, y1, x2, y2 = box
    for size in range(max_size, min_size - 1, -1):
        font = ImageFont.truetype(font_path, size)
        lines = text.split("\n")
        bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing, align="center")
        if bbox[2] - bbox[0] <= x2 - x1 - 24 and bbox[3] - bbox[1] <= y2 - y1 - 20:
            tx = (x1 + x2 - (bbox[2] - bbox[0])) / 2
            ty = (y1 + y2 - (bbox[3] - bbox[1])) / 2
            draw.multiline_text((tx, ty), text, font=font, fill="#17365D", spacing=spacing, align="center")
            return


def arrow(draw, start, end, color="#5B6573", width=5):
    draw.line([start, end], fill=color, width=width)
    import math
    ang = math.atan2(end[1]-start[1], end[0]-start[0])
    l = 16
    for d in (2.55, -2.55):
        p = (end[0] + l*math.cos(ang+d), end[1] + l*math.sin(ang+d))
        draw.line([end, p], fill=color, width=width)


def create_architecture_png():
    img = Image.new("RGB", (1800, 1260), "white")
    d = ImageDraw.Draw(img)
    title_font = ImageFont.truetype(FONT_CN_BOLD, 42)
    d.text((900, 36), "自主 MES：多模型、多 Agent 与可信执行架构", font=title_font, fill="#17365D", anchor="ma")
    layers = [
        (110, 130, 1690, 270, "目标与事件层", "经营目标 · 订单 · 插单 · 停机 · 缺料 · 质量异常 · 人员变化", "#EAF2F8", "#2E74B5"),
        (110, 315, 1690, 470, "Agent 协同层", "生产指挥 Agent · 排产 Agent · 工艺 Agent · 质量 Agent\n设备 Agent · 刀具 Agent · 物料 Agent · 维护 Agent", "#E9F5F5", "#2F8F9D"),
        (110, 515, 1690, 705, "混合模型层", "多模态大模型：理解、规划、解释、工具调用\n专业模型：视觉 · 时序 · 表格 · 因果 · 知识图谱 · 物理融合", "#F4F0FA", "#7156A5"),
        (110, 750, 1690, 915, "决策验证层", "CP-SAT / MILP 排产优化 · 离散事件仿真 · 数字孪生\n方案评估：交期 · 成本 · 质量 · 产能 · 风险", "#FFF4D6", "#C58A00"),
        (110, 960, 1690, 1120, "治理与执行层", "策略引擎 · 权限 · 审批 · 审计 · 回滚\nMES 工具 API → ERP / WMS / QMS / PLM → 边缘网关 → SCADA / PLC / CNC", "#F2F4F7", "#5B6573"),
    ]
    for x1, y1, x2, y2, label, text, fill, stroke in layers:
        rounded(d, (x1, y1, x2, y2), fill, stroke, 24, 4)
        d.rounded_rectangle((x1+18, y1+18, x1+250, y2-18), radius=16, fill=stroke)
        lf = ImageFont.truetype(FONT_CN_BOLD, 28)
        d.text((x1+134, (y1+y2)/2), label, font=lf, fill="white", anchor="mm")
        fit_text(d, text, (x1+280, y1+8, x2-20, y2-8), max_size=28, min_size=20)
    for y in (270, 470, 705, 915):
        arrow(d, (900, y+7), (900, y+38), "#7C8796", 6)
    foot = ImageFont.truetype(FONT_CN, 24)
    d.text((900, 1195), "原则：大模型负责协调；专业模型负责判断；优化器负责求解；控制系统负责确定性执行。", font=foot, fill="#17365D", anchor="mm")
    path = ASSET / "architecture.png"
    img.save(path, quality=95)
    return path


def create_closed_loop_png():
    img = Image.new("RGB", (1800, 820), "white")
    d = ImageDraw.Draw(img)
    title_font = ImageFont.truetype(FONT_CN_BOLD, 40)
    d.text((900, 34), "自主生产闭环", font=title_font, fill="#17365D", anchor="ma")
    labels = [
        ("1 感知", "采集订单、设备、\n刀具、质量和物料事件"),
        ("2 诊断", "识别异常并关联\n受影响工单与产品"),
        ("3 预测", "计算交期、质量、\n故障和产能风险"),
        ("4 规划", "生成重排、检验、\n维护和资源方案"),
        ("5 仿真", "验证约束并比较\n成本、交期和风险"),
        ("6 执行", "低风险自动执行；\n高风险请求审批"),
        ("7 学习", "记录效果并更新\n规则、模型和案例"),
    ]
    xs = [60, 305, 550, 795, 1040, 1285, 1530]
    fills = ["#EAF2F8", "#E9F5F5", "#F4F0FA", "#FFF4D6", "#EAF2F8", "#E9F5F5", "#F2F4F7"]
    strokes = ["#2E74B5", "#2F8F9D", "#7156A5", "#C58A00", "#2E74B5", "#287D5A", "#5B6573"]
    for i, ((head, body), x) in enumerate(zip(labels, xs)):
        rounded(d, (x, 230, x+205, 520), fills[i], strokes[i], 26, 4)
        hf = ImageFont.truetype(FONT_CN_BOLD, 27)
        d.text((x+102, 280), head, font=hf, fill=strokes[i], anchor="mm")
        fit_text(d, body, (x+10, 330, x+195, 490), max_size=23, min_size=18)
        if i < len(xs)-1:
            arrow(d, (x+210, 375), (xs[i+1]-8, 375), "#7C8796", 5)
    arrow(d, (1632, 550), (1632, 655), "#7C8796", 5)
    d.line([(1632,655),(160,655),(160,545)], fill="#7C8796", width=5)
    arrow(d, (160,600), (160,535), "#7C8796", 5)
    ff = ImageFont.truetype(FONT_CN, 24)
    d.text((900, 710), "闭环的价值不在于“给建议”，而在于建议能够被验证、授权、执行、回滚并持续改进。", font=ff, fill="#17365D", anchor="mm")
    path = ASSET / "closed-loop.png"
    img.save(path, quality=95)
    return path


def create_roadmap_png():
    img = Image.new("RGB", (1800, 900), "white")
    d = ImageDraw.Draw(img)
    title_font = ImageFont.truetype(FONT_CN_BOLD, 40)
    d.text((900, 32), "分阶段自治路线", font=title_font, fill="#17365D", anchor="ma")
    stages = [
        ("阶段 1｜可信数据底座", "0–6个月", "MES核心闭环\n设备与事件连接\n制造语义模型\n知识库与自然语言查询", "#EAF2F8", "#2E74B5"),
        ("阶段 2｜预测辅助", "6–12个月", "工时和交期预测\n质量风险预警\n刀具与设备异常\n图纸和工艺辅助", "#E9F5F5", "#2F8F9D"),
        ("阶段 3｜L3自主闭环", "12–24个月", "动态排产\n质量防扩散\n维护协调\n低风险自动执行", "#FFF4D6", "#C58A00"),
        ("阶段 4｜局部L4自治", "24–36个月", "单车间目标驱动\n数字孪生持续优化\n有限自适应工艺\n人类管理例外", "#F4F0FA", "#7156A5"),
    ]
    xvals = [75, 500, 925, 1350]
    for i, (head, period, body, fill, stroke) in enumerate(stages):
        x = xvals[i]
        rounded(d, (x, 190, x+350, 690), fill, stroke, 28, 4)
        hf = ImageFont.truetype(FONT_CN_BOLD, 27)
        pf = ImageFont.truetype(FONT_CN_BOLD, 23)
        d.text((x+175, 250), head, font=hf, fill=stroke, anchor="mm")
        d.text((x+175, 315), period, font=pf, fill="#5B6573", anchor="mm")
        fit_text(d, body, (x+25, 365, x+325, 645), max_size=25, min_size=19)
        if i < 3:
            arrow(d, (x+355, 440), (xvals[i+1]-8, 440), "#7C8796", 6)
    ff = ImageFont.truetype(FONT_CN, 23)
    d.text((900, 780), "推荐目标：首版全系统达到 L2，并让排产、质量、维护、跟催四个闭环逐步达到 L3。", font=ff, fill="#17365D", anchor="mm")
    path = ASSET / "roadmap.png"
    img.save(path, quality=95)
    return path


def add_picture(doc, path, width=6.45, caption=None, alt=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    if alt:
        doc_pr = shape._inline.docPr
        doc_pr.set("descr", alt)
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(8)
        r = cp.add_run(caption)
        set_run_font(r, 9, color=GRAY)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15
for name, size, color, before, after in [
    ("Title", 28, NAVY, 0, 8),
    ("Subtitle", 14, GRAY, 0, 16),
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 11.5, NAVY, 8, 4),
]:
    st = styles[name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.color.rgb = rgb(color)
    st.font.bold = name != "Subtitle"
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

# Running header/footer.
hp = sec.header.paragraphs[0]
hp.text = "自主 MES 产品规划与流程架构"
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_run_font(hp.runs[0], 9, color=GRAY)
fp = sec.footer.paragraphs[0]
add_page_number(fp)

# Cover: editorial cover, restrained and technical.
doc.add_paragraph().paragraph_format.space_after = Pt(70)
k = doc.add_paragraph()
k.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = k.add_run("产品战略与系统蓝图")
set_run_font(r, 11, True, GOLD)
k.paragraph_format.space_after = Pt(18)
t = doc.add_paragraph(style="Title")
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("自主 MES 产品规划与流程架构")
set_run_font(r, 28, True, NAVY)
s = doc.add_paragraph(style="Subtitle")
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run("面向机械加工与装备制造的多模型、多 Agent 自主生产运营平台")
set_run_font(r, 14, False, GRAY)
doc.add_paragraph().paragraph_format.space_after = Pt(40)
add_callout(doc, "北极星定义", "以 MES 为可信执行底座，以数字孪生为验证环境，以多 Agent 为生产组织者，以策略、权限和人类审批保障安全，让系统从记录生产升级为受控地组织生产。", PALE_BLUE, BLUE)
doc.add_paragraph().paragraph_format.space_after = Pt(32)
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run(f"规划版本 V1.0  |  {date.today().isoformat()}  |  概念规划稿")
set_run_font(r, 10, color=GRAY)
doc.add_page_break()

doc.add_heading("文档导读", level=1)
add_para(doc, "本文件用于统一产品愿景、技术架构、业务闭环、实施路径和商业预期，既可作为产品立项材料，也可作为后续 PRD、技术方案、试点方案和融资材料的上位蓝图。")
add_callout(doc, "核心判断", "本产品具有形成新一代工业软件范式的潜力，但革命性不来自“MES + AI”，而来自工厂能否安全地把一部分日常生产决策权交给系统，并持续获得可量化收益。", PALE_GOLD, GOLD)
doc.add_heading("目录", level=2)
toc = [
    "1. 战略定位与产品愿景", "2. 目标客户、问题与价值", "3. 自主生产业务闭环",
    "4. 总体系统与流程架构", "5. 模型组合与模型路由", "6. 多 Agent 组织结构",
    "7. 数据、数字孪生与集成底座", "8. 自主权、安全与治理", "9. 产品功能规划",
    "10. 分阶段实施路线", "11. 可实现度评估", "12. 收益预判与 ROI",
    "13. 试点设计与验收指标", "14. 核心风险与应对", "15. 下一步工作计划"
]
for item in toc:
    add_bullet(doc, item)

doc.add_page_break()
doc.add_heading("1. 战略定位与产品愿景", level=1)
doc.add_heading("1.1 产品定义", level=2)
add_para(doc, "产品暂定定位：面向多品种、小批量机械加工与装备制造的自主生产运营平台（Agentic / Autonomous MES）。")
add_para(doc, "它不是在传统 MES 页面上增加聊天入口，而是把生产目标、制造事件、专业模型、优化求解、数字孪生和受控执行组织成持续运行的决策闭环。")
doc.add_heading("1.2 范式变化", level=2)
add_table(doc, ["传统 MES", "自主 MES"], [
    ("人制定计划，系统记录结果", "人定义目标与边界，系统持续组织生产"),
    ("依赖菜单和人工操作", "以事件触发 Agent，自主调用工具"),
    ("报表回答发生了什么", "系统进一步诊断原因、预测后果和执行方案"),
    ("异常由人发现和协调", "系统主动发现、仿真、处置并验证效果"),
    ("经验存在于个人", "经验沉淀为知识、策略、模型和决策轨迹"),
], [4680, 4680])
doc.add_heading("1.3 革命性成立条件", level=2)
for x in [
    "建议必须进入可执行闭环，而不只是生成报告。",
    "每项动作必须可解释、可授权、可审计和可回滚。",
    "不同工厂部署必须逐步标准化，不能完全依赖驻场定制。",
    "必须通过准交率、吞吐、停机、质量、WIP和人工效率证明经济价值。",
]: add_bullet(doc, x)

doc.add_heading("2. 目标客户、问题与价值", level=1)
doc.add_heading("2.1 优先客户画像", level=2)
add_bullet(doc, "机械加工、精密零件、模具、工装、航空航天零部件及复杂装备制造企业。")
add_bullet(doc, "多品种、小批量、订单变化快，存在频繁插单、返工、设备故障和物料波动。")
add_bullet(doc, "拥有 10 台以上数控设备，工艺路线、刀具夹具和人员资质约束明显。")
add_bullet(doc, "已有 ERP 或基础 MES，但现场数据割裂、计划失真、异常处理依赖个人。")
doc.add_heading("2.2 关键业务问题", level=2)
add_table(doc, ["问题", "传统方式", "自主 MES 目标"], [
    ("工时不准", "按标准工时或经验估算", "基于历史、几何、设备和人员动态预测"),
    ("插单与停机", "计划员手工反复调整", "数分钟内生成并仿真重排方案"),
    ("质量异常", "检验后再调查", "预测风险、冻结影响范围并追加检验"),
    ("设备与刀具", "定期维护、固定寿命", "基于状态和任务动态安排维护与换刀"),
    ("工艺知识", "依赖老师傅和分散文档", "图纸解析、相似案例检索、经验复用"),
], [1800, 3000, 4560], font_size=9.2)
doc.add_heading("2.3 北极星价值", level=2)
add_callout(doc, "核心价值主张", "任何插单、停机、缺料、返工或质量风险发生后，系统在数分钟内给出经过约束检查和仿真的调整方案；低风险场景自动执行，高风险场景只需管理者批准。", PALE_CYAN, CYAN)

doc.add_page_break()
doc.add_heading("3. 自主生产业务闭环", level=1)
add_picture(doc, create_closed_loop_png(), caption="图 1  自主生产从感知到学习的闭环", alt="自主生产闭环：感知、诊断、预测、规划、仿真、执行、学习")
doc.add_heading("3.1 示例：关键机床异常后的自动调控", level=2)
steps = [
    "设备 Agent 识别主轴振动、负载和温升的异常组合，并判断风险等级。",
    "质量 Agent 追溯异常时间窗内的工单、序列号、刀具和关键尺寸，自动冻结潜在影响产品。",
    "维护 Agent 创建维护任务，检查人员、备件和最小影响维护窗口。",
    "排产 Agent 寻找可替代设备，验证精度、行程、程序、刀具夹具和人员资质。",
    "优化器与数字孪生比较交期、成本、换型、加班和质量风险。",
    "策略引擎根据影响范围决定自动执行或提交审批。",
    "MES 工具完成暂停、转移、追加检验、维修派工和通知，并持续验证结果。",
]
for x in steps: add_number(doc, x)
doc.add_heading("3.2 实时控制边界", level=2)
add_callout(doc, "安全边界", "Agent 可以调控工单、队列、检验、维护、物料和资源，但不得绕过 PLC/CNC 安全联锁。毫秒级运动控制、轴动作和安全停机继续由确定性控制系统负责。", "FCECEC", RED)

doc.add_page_break()
doc.add_heading("4. 总体系统与流程架构", level=1)
add_picture(doc, create_architecture_png(), caption="图 2  自主 MES 总体架构", alt="目标事件、Agent、混合模型、决策验证、治理执行五层架构")
doc.add_heading("4.1 架构原则", level=2)
for x in [
    "事件驱动：订单、设备、质量和物料变化主动触发 Agent。",
    "混合智能：大模型负责理解与协调，专业模型负责预测，优化器负责求解。",
    "工具化执行：Agent 只能通过受控 MES API 行动，不直接修改数据库。",
    "仿真优先：高影响决策先在数字孪生中验证，再进入审批或执行。",
    "边云协同：低延迟视觉和时序模型部署在边缘，跨工厂模型和知识在中心运行。",
]: add_bullet(doc, x)

doc.add_page_break()
doc.add_heading("5. 模型组合与模型路由", level=1)
add_para(doc, "不存在一种模型可以完成全部工业任务。产品需要模型路由器根据输入模态、任务类型、时效、风险、成本和数据位置选择最合适的模型。")
add_table(doc, ["能力", "主要模型/算法", "输出与用途"], [
    ("Agent中枢", "多模态大模型、RAG、工具调用", "理解目标、任务拆解、调用模型和工具、解释方案"),
    ("图纸与工艺", "VLM + OCR + CAD/STEP解析", "尺寸公差、几何特征、工艺要求、相似零件"),
    ("工时与交期", "LightGBM、CatBoost、表格Transformer", "实际工时、完工时间和延期概率"),
    ("设备与刀具", "信号处理 + TCN/Transformer/Autoencoder", "异常评分、健康状态和剩余寿命"),
    ("视觉质量", "目标检测、分割、视觉异常检测", "缺陷位置、类型和风险"),
    ("根因分析", "因果图、贝叶斯网络、过程挖掘", "相关因素、因果假设和验证路径"),
    ("动态排产", "CP-SAT、MILP、启发式、多目标优化", "满足硬约束的可执行排程"),
    ("工艺优化", "机理模型、MPC、贝叶斯优化", "在验证边界内推荐或调整参数"),
], [1900, 3260, 4200], font_size=8.9)
doc.add_heading("5.1 模型选择原则", level=2)
add_bullet(doc, "先用成熟通用模型与专业小模型验证价值，不从零训练工业大模型。")
add_bullet(doc, "中小规模结构化数据优先考虑可解释、稳定的树模型。")
add_bullet(doc, "高频设备信号先经过物理与信号处理，再进入时序模型。")
add_bullet(doc, "排产由优化器求解，大模型只负责理解目标、生成约束和解释结果。")
add_bullet(doc, "模型输出必须包含置信度、数据范围、版本和适用边界。")

doc.add_page_break()
doc.add_heading("6. 多 Agent 组织结构", level=1)
add_table(doc, ["Agent", "职责", "可调用工具", "初始自治"], [
    ("生产指挥", "协调交期、成本、质量和产能目标", "查询全局状态、发起协同、升级冲突", "L2"),
    ("排产", "插单、停机、返工后的动态重排", "运行优化、调整队列、转移工序", "L2→L3"),
    ("工艺", "图纸解析、路线和资源推荐", "检索案例、生成草案、提交审核", "L1→L2"),
    ("质量", "风险预测、防扩散、根因分析", "冻结产品、追加检验、创建NCR", "L3"),
    ("设备", "健康评估、能力和负荷管理", "限制排产、发出预警、请求停机", "L2→L3"),
    ("刀具", "寿命、齐套和换刀窗口", "预留刀具、生成备刀、限制任务", "L3"),
    ("物料", "齐套、缺料预测与配送", "预留物料、请求配送、异常升级", "L3"),
    ("维护", "维护计划、人员和备件协调", "创建维护工单、安排窗口", "L3"),
], [1300, 2920, 3380, 1760], font_size=8.6)
doc.add_heading("6.1 Agent协同规则", level=2)
for x in [
    "生产指挥 Agent 负责多目标冲突，不替代专业 Agent 的领域判断。",
    "硬约束由规则和优化器执行，不能由大模型自由解释。",
    "多个 Agent 对同一资源产生冲突时，进入仲裁流程并保留备选方案。",
    "Agent 的权限按工厂、车间、设备、动作类型和风险等级最小化授予。",
]: add_bullet(doc, x)

doc.add_page_break()
doc.add_heading("7. 数据、数字孪生与集成底座", level=1)
doc.add_heading("7.1 统一制造上下文", level=2)
add_para(doc, "AI 必须能够把某一时刻的设备信号关联到具体设备、工单、工序、产品序列号、材料批次、刀具、程序版本、人员和最终质量结果。没有这些上下文，传感器数据很难形成可执行价值。")
add_table(doc, ["数据域", "关键对象", "核心要求"], [
    ("产品与工艺", "产品、BOM、工艺路线、工序、图纸、程序", "版本可追溯，变更可关联"),
    ("生产执行", "订单、工单、派工、报工、WIP、返工", "事件时间准确，状态一致"),
    ("资源", "设备、人员、刀具、夹具、量具", "能力、状态、资质和寿命结构化"),
    ("质量", "检验计划、测量值、缺陷、不合格处置", "与产品和过程参数关联"),
    ("设备时序", "电流、振动、温度、负载、报警", "边缘过滤、时间同步、语义化"),
    ("决策轨迹", "触发、方案、审批、执行、效果", "可审计并可用于持续学习"),
], [1700, 3580, 4080], font_size=9)
doc.add_heading("7.2 集成范围", level=2)
add_bullet(doc, "向上：ERP、PLM、APS、WMS、QMS、CMMS/EAM。")
add_bullet(doc, "向下：SCADA、PLC、CNC、机器人、测量设备、视觉系统。")
add_bullet(doc, "协议与通道：REST/gRPC、消息总线、OPC UA、MQTT、Modbus及设备厂商接口。")

doc.add_page_break()
doc.add_heading("8. 自主权、安全与治理", level=1)
add_table(doc, ["等级", "能力", "人员角色", "近期策略"], [
    ("L0", "传统MES，仅记录和展示", "人完成全部判断", "兼容"),
    ("L1", "AI查询、解释和总结", "人判断并执行", "首版覆盖"),
    ("L2", "Agent发现问题并生成方案", "人批准与执行", "首版目标"),
    ("L3", "低风险动作自动执行", "人处理高风险例外", "四大闭环目标"),
    ("L4", "单车间范围自主协调", "人管理目标和例外", "中长期试点"),
    ("L5", "跨工厂高度自治", "人负责战略治理", "非近期承诺"),
], [900, 3650, 3000, 1810], font_size=9.2)
doc.add_heading("8.1 动作权限示例", level=2)
add_table(doc, ["动作", "建议自治", "约束"], [
    ("生成日报、异常摘要", "L3自动", "保留数据来源"),
    ("追加质量检验", "L3自动", "不减少法规/客户要求"),
    ("创建维护工单", "L3自动", "不直接触发危险动作"),
    ("调整同设备任务队列", "L3自动", "不改变工艺路线和交期承诺"),
    ("跨设备重排", "L2审批", "验证设备能力、程序、刀夹具和人员"),
    ("修改关键工艺参数", "L1建议", "工艺工程师审批和验证"),
    ("绕过安全联锁", "禁止", "任何情况下均不可授权"),
], [2700, 1780, 4880], font_size=9)
doc.add_page_break()
doc.add_heading("8.2 决策审计记录", level=2)
add_para(doc, "每次自主行动应保存：触发事件、业务目标、证据、模型及版本、候选方案、选中原因、风险等级、授权策略、执行动作、回滚方式、预期收益和实际效果。")

doc.add_heading("9. 产品功能规划", level=1)
add_table(doc, ["产品域", "基础能力", "AI/Agent增强"], [
    ("生产计划", "工单、工序、派工、进度", "工时预测、交期风险、动态重排"),
    ("工艺管理", "BOM、路线、作业指导、版本", "图纸解析、相似工艺、路线草案"),
    ("现场执行", "扫码、开工、报工、异常", "语音交互、主动跟催、异常结构化"),
    ("质量", "检验、SPC、NCR、追溯", "风险预测、防扩散、因果根因"),
    ("设备与刀具", "状态、点检、寿命、维护", "异常检测、剩余寿命、维护窗口"),
    ("物料与WIP", "批次、齐套、配送、在制品", "缺料预测、动态配送和WIP优化"),
    ("运营驾驶舱", "OEE、良率、准交率、产能", "目标驱动、方案比较、收益核算"),
], [1750, 3580, 4030], font_size=9)

doc.add_page_break()
doc.add_heading("10. 分阶段实施路线", level=1)
add_picture(doc, create_roadmap_png(), caption="图 3  从可信数据底座到局部 L4 自治", alt="四阶段自主MES实施路线")
doc.add_heading("10.1 推荐首期范围", level=2)
add_bullet(doc, "选择一个约 10–30 台关键设备的机械加工车间。")
add_bullet(doc, "聚焦实际工时、动态排产、质量防扩散、维护协调和生产跟催。")
add_bullet(doc, "首期不承诺自动修改 CNC 程序或关键加工参数。")
add_bullet(doc, "先完成影子运行，再逐步开启低风险自动执行。")

doc.add_heading("11. 可实现度评估", level=1)
add_table(doc, ["能力阶段", "实现度", "结论"], [
    ("AI知识助手与自然语言MES", "95%", "技术成熟，可优先实现"),
    ("图纸、工艺与现场多模态助手", "80–90%", "关键数据仍需人工复核"),
    ("工时、交期和质量风险预测", "80–90%", "数据质量决定上限"),
    ("动态排产和插单仿真", "约85%", "算法成熟，约束建模是难点"),
    ("设备和刀具预测", "70–85%", "依赖传感器和有效标签"),
    ("低风险Agent自动执行", "75–85%", "依赖工具、权限和回滚"),
    ("单车间有限自治L4", "55–70%", "需长期运行和安全验证"),
    ("全工厂无人自治L5", "15–30%", "不适合作为近期承诺"),
], [3560, 1600, 4200], font_size=9.1)
add_callout(doc, "综合判断", "以 L2 全覆盖、四个核心闭环达到 L3 为目标，当前综合可实现度约 80%。若要求直接自主控制 CNC 参数并实现全工厂无人决策，可实现度下降至约 30%–45%。", PALE_GOLD, GOLD)

doc.add_page_break()
doc.add_heading("12. 收益预判与 ROI", level=1)
doc.add_heading("12.1 成熟运行后的经营改善区间", level=2)
add_table(doc, ["指标", "保守", "基准", "理想"], [
    ("准时交付率", "+3个百分点", "+8个百分点", "+15个百分点"),
    ("有效产能/吞吐", "+5%", "+12%", "+20%"),
    ("非计划停机", "-10%", "-25%", "-40%"),
    ("换型和等待", "-10%", "-25%", "-40%"),
    ("报废与返工", "-8%", "-20%", "-35%"),
    ("在制品WIP", "-8%", "-20%", "-35%"),
    ("计划员重复工作", "-30%", "-60%", "-80%"),
], [2800, 2180, 2180, 2200], aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])
doc.add_page_break()
doc.add_heading("12.2 亿元级工厂基准算例", level=2)
add_para(doc, "假设年营业收入 1 亿元、制造成本 6500 万元、30–50 台数控设备、订单需求充足。在避免重复计算的前提下，成熟运行后的年度可兑现收益预计约 470 万–760 万元，中位数约 500 万–700 万元。")
add_table(doc, ["收益来源", "年度预估"], [
    ("有效产能增加贡献利润", "200万–300万元"),
    ("报废返工减少", "50万–80万元"),
    ("停机与异常损失减少", "100万–150万元"),
    ("计划、跟催和统计效率", "50万–80万元"),
    ("WIP资金和持有成本", "20万–50万元"),
    ("工艺准备与知识复用", "50万–100万元"),
], [6000, 3360], font_size=9.5)
add_callout(doc, "投资判断", "中型工厂若首年投入约 200万–500万元，基准回收期可按 12–24个月测算；三年净 ROI 可设为 80%–200% 的验证区间。所有收益均需通过试点基线数据重新测算。", PALE_CYAN, CYAN)

doc.add_heading("13. 试点设计与验收指标", level=1)
doc.add_heading("13.1 试点闭环", level=2)
for x in [
    "动态排产：发生插单、停机或返工后，在 5 分钟内生成可执行方案。",
    "质量防扩散：识别风险窗口，自动关联并冻结受影响序列号，生成追加检验。",
    "维护协调：根据设备风险生成维护任务，并给出对交期影响最小的窗口。",
    "生产跟催：自动识别潜在延期、缺料和工序拥堵，形成处置任务。",
]: add_bullet(doc, x)
doc.add_page_break()
doc.add_heading("13.2 建议验收指标", level=2)
add_table(doc, ["类别", "指标", "建议目标"], [
    ("数据", "工单-工序-设备-产品关联完整率", "≥95%"),
    ("排产", "方案生成时间", "≤5分钟"),
    ("执行", "建议可执行率", "≥80%"),
    ("采用", "计划员采纳或少量修改比例", "≥70%"),
    ("交付", "准交率改善", "+5个百分点以上"),
    ("异常", "平均响应时间", "下降50%以上"),
    ("质量", "风险批次防扩散覆盖率", "≥90%"),
    ("安全", "未授权高风险动作", "0次"),
], [1500, 4920, 2940], font_size=9.2)

doc.add_page_break()
doc.add_heading("14. 核心风险与应对", level=1)
add_table(doc, ["风险", "表现", "应对策略"], [
    ("数据不可用", "时间戳、标识、原因码和质量关联缺失", "先建设最小制造事件模型和数据质量看板"),
    ("模型幻觉", "大模型生成不存在的工艺或原因", "RAG、结构化工具、规则校验和来源显示"),
    ("建议不可执行", "忽略夹具、资质、程序或物料约束", "硬约束进入优化器和策略引擎"),
    ("过度自治", "AI动作影响安全、质量或交付承诺", "分级授权、影子运行、审批、回滚和熔断"),
    ("部署定制化", "每个工厂都需要从头开发", "标准语义模型、连接器、Agent工具和行业模板"),
    ("收益难证明", "有预测准确率但无经营改善", "记录建议-执行-效果轨迹，持续核算价值"),
    ("人员抵触", "计划员和工程师不信任系统", "先辅助后自治，展示依据，保留人工否决权"),
], [1800, 3100, 4460], font_size=8.9)

doc.add_page_break()
doc.add_heading("15｜下一步工作计划", level=1)
for x in [
    "确定首个细分场景：精密机加工、模具、非标装备或航空零部件。",
    "选择试点工厂并完成两周业务发现，建立收益基线。",
    "输出领域模型：核心对象、事件、状态机、约束和Agent工具清单。",
    "完成首版技术验证：工时预测 + 动态排产 + 事件驱动Agent。",
    "建设影子运行环境，以历史数据回放和实时旁路方式验证。",
    "制定L2到L3授权策略，选择可回滚的低风险动作。",
    "运行8–12周试点，依据KPI决定扩大设备、车间和Agent范围。",
]: add_number(doc, x, style="List Number 2")
add_callout(doc, "建议的第一产品楔子", "优先做“动态生产调度与异常自治”：任何插单、停机、缺料和返工发生后，自动生成经过仿真的重排方案，并在授权范围内执行。该能力可自然连接质量、维护、物料和工艺，是最有机会形成平台效应的入口。", PALE_GOLD, GOLD)

doc.add_page_break()
doc.add_heading("附录 A：关键参考与边界说明", level=1)
refs = [
    "ISA-95：企业系统与制造控制系统集成框架，https://www.isa.org/standards-and-publications/isa-standards/isa-95-standard",
    "NIST 2026 AI/ML for Smart Manufacturing Roadmap，https://www.nist.gov/publications/2026-roadmap-artificial-intelligence-and-machine-learning-smart-manufacturing",
    "NIST Digital Twins for Advanced Manufacturing，https://www.nist.gov/digital-twins",
    "NIST Augmented Intelligence for Manufacturing Systems，https://www.nist.gov/programs-projects/augmented-intelligence-manufacturing-systems-aims",
    "Microsoft Industrial Foundation Models，https://www.microsoft.com/en-us/research/articles/towards-industrial-foundation-models-integrating-large-language-models-with-industrial-data-intelligence/",
    "NIST Economics of Digital Twins，https://www.nist.gov/publications/economics-digital-twins-costs-benefits-and-economic-decision-making",
]
for x in refs: add_bullet(doc, x)
add_para(doc, "说明：可实现度、收益和ROI均为概念规划阶段的区间预判，不构成对具体工厂的承诺。正式商业测算必须基于设备、订单、成本、质量、人员和停机基线进行。", color=GRAY, italic=True)

# Document metadata.
doc.core_properties.title = "自主 MES 产品规划与流程架构"
doc.core_properties.subject = "机械加工与装备制造的多模型、多Agent自主生产运营平台"
doc.core_properties.author = "项目规划组"
doc.core_properties.keywords = "MES, AI Agent, 机械加工, 装备制造, 数字孪生, 动态排产"

doc.save(OUT)
print(OUT)
